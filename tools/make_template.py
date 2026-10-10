"""
One-off script: turn the CPRI report format (reference/Transformer_Test_report_format_1.docx)
into a docxtpl/Jinja template (templates/report_template.docx) WITHOUT changing its formatting.

    python tools/make_template.py

What it does (python-docx, tables addressed by index - see docs/architecture.md):
  1. fills empty value cells / XXXX placeholders with Jinja tags ({{ report_no }}, {{ sp.rated_power }} ...)
     keeping the run formatting of the original cell
  2. turns repeating rows into {%tr for %} loops (voltage ratio taps, drawings)
  3. wraps optional rows ({%tr if %}) and sections/sheets ({%p if %}) so only requested tests appear
  4. replaces "N of 11" sheet references with computed sheet numbers
  5. makes every sheet start on a new page (explicit page break) so removing a
     sheet never lets two sheets run together
  6. puts the seal hash into an empty footer cell and a QR code on the cover
"""
import copy
import os
import re
import sys

from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, "reference", "Transformer_Test_report_format_1.docx")
DST = os.path.join(ROOT, "templates", "report_template.docx")


# ---------------------------------------------------------------- XML helpers
def own_t_nodes(p):
    """w:t nodes that belong to paragraph p itself (not to nested text-box paragraphs)."""
    out = []
    for t in p.iter(qn("w:t")):
        anc = next(t.iterancestors(qn("w:p")), None)
        if anc is p:
            out.append(t)
    return out


def p_text(p):
    return "".join(t.text or "" for t in own_t_nodes(p))


def regex_replace_p(p, pattern, repl):
    """Regex replace across the runs of one paragraph, keeping the formatting of the first run of each match."""
    nodes = own_t_nodes(p)
    if not nodes:
        return 0
    full = "".join(t.text or "" for t in nodes)
    matches = list(re.finditer(pattern, full))
    if not matches:
        return 0
    for m in reversed(matches):
        new = m.expand(repl) if isinstance(repl, str) else repl(m)
        start, end = m.start(), m.end()
        pos = 0
        first = True
        for t in nodes:
            txt = t.text or ""
            a, b = pos, pos + len(txt)
            pos = b
            if b <= start or a >= end:
                if a == b == start and first and start == end:
                    pass
                continue
            s0, s1 = max(start, a) - a, min(end, b) - a
            if first:
                t.text = txt[:s0] + new + txt[s1:]
                first = False
            else:
                t.text = txt[:s0] + txt[s1:]
            t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    return len(matches)


def replace_everywhere(root, pattern, repl):
    n = 0
    for p in list(root.iter(qn("w:p"))):
        n += regex_replace_p(p, pattern, repl)
    return n


def unique_cells(row):
    """python-docx returns merged cells repeatedly; keep one entry per w:tc."""
    out, seen = [], set()
    for c in row.cells:
        if id(c._tc) in seen:
            continue
        seen.add(id(c._tc))
        out.append(c)
    return out


def set_tc_text(tc, text):
    """Replace the whole content of a table cell with `text`, keeping the first run's formatting."""
    paras = tc.findall(qn("w:p"))
    p0 = paras[0]
    for extra in paras[1:]:
        tc.remove(extra)
    runs = p0.findall(qn("w:r"))
    rpr = None
    if runs:
        rpr = runs[0].find(qn("w:rPr"))
    else:
        ppr = p0.find(qn("w:pPr"))
        mark = ppr.find(qn("w:rPr")) if ppr is not None else None
        if mark is not None:
            rpr = copy.deepcopy(mark)
            rpr.tag = qn("w:rPr")
    for child in list(p0):
        if child.tag != qn("w:pPr"):
            p0.remove(child)
    if rpr is not None:
        # the format marks some "fill here" cells in red/bold/underline - values print in normal black
        color = rpr.find(qn("w:color"))
        if color is not None and color.get(qn("w:val"), "").upper() in ("FF0000", "C00000", "RED"):
            rpr = copy.deepcopy(rpr)
            for tag in ("w:color", "w:u", "w:b", "w:bCs"):
                for el in rpr.findall(qn(tag)):
                    rpr.remove(el)
    r = OxmlElement("w:r")
    if rpr is not None:
        r.append(copy.deepcopy(rpr))
    t = OxmlElement("w:t")
    t.text = text
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    r.append(t)
    p0.append(r)


def set_cell(table, r, c, text):
    set_tc_text(unique_cells(table.rows[r])[c]._tc, text)


def tag_row(table, row_idx, tag, where):
    """Insert a docxtpl {%tr ... %} row before/after table.rows[row_idx]."""
    tr = table.rows[row_idx]._tr
    new = copy.deepcopy(tr)
    tcs = new.findall(qn("w:tc"))
    for i, tc in enumerate(tcs):
        set_tc_text(tc, tag if i == 0 else "")
    if where == "before":
        tr.addprevious(new)
    else:
        tr.addnext(new)


def wrap_rows(table, first, last, cond):
    """{%tr if cond %} rows first..last {%tr endif %} (indices refer to the table before wrapping)."""
    rows = [table.rows[i]._tr for i in range(first, last + 1)]
    tag_row(table, last, "{%tr endif %}", "after")
    first_tr = rows[0]
    new = copy.deepcopy(first_tr)
    for i, tc in enumerate(new.findall(qn("w:tc"))):
        set_tc_text(tc, "{%tr if " + cond + " %}" if i == 0 else "")
    first_tr.addprevious(new)


def make_tag_p(text):
    p = OxmlElement("w:p")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = text
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    r.append(t)
    p.append(r)
    return p


def wrap_block(first_el, last_el, cond):
    first_el.addprevious(make_tag_p("{%p if " + cond + " %}"))
    last_el.addnext(make_tag_p("{%p endif %}"))


def is_blank(p):
    if p.tag != qn("w:p"):
        return False
    if p_text(p).strip():
        return False
    for bad in ("w:drawing", "w:pict", "w:object", "w:sectPr", "w:br"):
        if p.find(".//" + qn(bad)) is not None:
            return False
    if p.find(".//{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent") is not None:
        return False
    return True


def set_page_break_before(p):
    ppr = p.find(qn("w:pPr"))
    if ppr is None:
        ppr = OxmlElement("w:pPr")
        p.insert(0, ppr)
    pb = OxmlElement("w:pageBreakBefore")
    # pageBreakBefore must come early in pPr (after pStyle/keepNext/keepLines)
    anchor = None
    for tag in ("w:pStyle", "w:keepNext", "w:keepLines"):
        el = ppr.find(qn(tag))
        if el is not None:
            anchor = el
    if anchor is not None:
        anchor.addnext(pb)
    else:
        ppr.insert(0, pb)


# ---------------------------------------------------------------- main
def main():
    doc = Document(SRC)
    P = doc.paragraphs
    T = doc.tables
    Pe = [p._p for p in P]
    assert len(T) == 39, f"unexpected template: {len(T)} tables"

    # ======================= Sheet 1: cover (T0)
    t = T[0]
    set_cell(t, 2, 2, "{{ customer_block }}")
    set_cell(t, 4, 2, "{{ manufacturer_block }}")
    set_cell(t, 6, 2, "{{ particulars }}")
    set_cell(t, 8, 2, "{{ sp.type }}")
    set_cell(t, 10, 2, "{{ sp.serial }}")
    set_cell(t, 11, 2, "{{ n_samples_words }}")
    set_cell(t, 13, 2, "{{ test_dates }}")
    set_cell(t, 15, 2, "{{ sample_code }}")
    # r17: first paragraph = particulars of tests, second = (Refer Sheet 3 of 11) handled by the generic pass
    tc = unique_cells(t.rows[17])[2]._tc
    first_p = tc.findall(qn("w:p"))[0]
    regex_replace_p(first_p, r"^.*$", "{{ tests_particulars }}")
    set_cell(t, 19, 2, "{{ standard }}")
    set_cell(t, 23, 2, "{{ customer_requirement }}")
    set_cell(t, 24, 2, "{{ deviations }}")
    set_cell(t, 27, 2, "{{ witness_customer }}")
    set_cell(t, 28, 2, "{{ witness_other }}")
    for r, key in zip(range(33, 39), ("n_sheets_words", "n_osc_words", "n_graph_words", "n_photo_words",
                                      "n_circuit_words", "n_drawing_words")):
        set_cell(t, r, 2, "{{ " + key + " }}")
    set_cell(t, 41, 2, "{{ qr }}")
    # reviewer signature (right) - "XXXXXX"
    sig = unique_cells(t.rows[43])
    regex_replace_p(sig[-1]._tc.findall(qn("w:p"))[0], r"X{6}", "{{ reviewer_name }}")

    # ======================= Sheet 2: description of sample (T2)
    t = T[2]
    desc = ["sp.test_sample", "sp.type", "sp.serial", "sp.rated_power", "sp.rated_voltage", "sp.um", None,
            "sp.hv_ins", "sp.lv_ins", "sp.impedance", "sp.phases", "sp.freq", "sp.vector_group", "sp.tappings",
            "sp.eel", "sp.cooling", "sp.max_tl50", "sp.max_tl100", "sp.coil", "sp.core", "sp.mfg", "sp.oil"]
    for r, key in enumerate(desc):
        if key:
            set_cell(t, r, 2, "{{ " + key + " }}")
    set_cell(t, 24, 2, "{{ sp.lv_winding }}")
    set_cell(t, 25, 2, "{{ sp.hv_winding }}")

    # ======================= Sheet 3: summary (T4, T5, T6)
    set_cell(T[4], 1, 2, "Refer Sheet {{ first_test_sheet }} of {{ n_sheets }} to Sheet {{ last_test_sheet }} of {{ n_sheets }}")
    t = T[5]
    set_cell(t, 7, 0, "{{ routine_heading }}")
    # rows: 1 special hdr, 2 sc, 3 nlc112, 4 type hdr, 5 temp, 6 ptype, 7 routine hdr, 8..16 routine tests
    row_conds = {1: "g.special", 2: "t.sc_withstand", 3: "t.nlc_112", 4: "g.type", 5: "t.temp_rise",
                 6: "t.pressure_type", 7: "g.routine", 8: "t.winding_res", 9: "t.ratio", 10: "t.impedance",
                 11: "t.nll", 12: "t.ir", 13: "t.induced", 14: "t.sep_source", 15: "t.pressure_routine", 16: "t.oil_leak"}
    for r in sorted(row_conds, reverse=True):     # bottom-up so indices stay valid
        wrap_rows(t, r, r, row_conds[r])
    t = T[6]
    set_cell(t, 0, 2, "{{ osc_list }}")
    set_cell(t, 2, 2, "{{ graph_list }}")
    set_cell(t, 4, 2, "{{ photo_list }}")
    set_cell(t, 6, 2, "{{ circuit_list }}")

    # ======================= Sheet 4: drawings (T8) -> loop
    t = T[8]
    for c, k in enumerate(("d.sl", "d.number", "d.sheet", "d.rev")):
        set_cell(t, 1, c, "{{ " + k + " }}")
    for _ in range(3):
        t._tbl.remove(t.rows[2]._tr)
    tag_row(t, 1, "{%tr for d in drawings %}", "before")
    tag_row(t, 2, "{%tr endfor %}", "after")

    # ======================= Sheet 5: routine tests
    t = T[11]
    for c, k in enumerate(("bt.n", "bt.h", "bt.l", "at.n", "at.h", "at.l"), start=1):
        set_cell(t, 4, c, "{{ wr.hv." + k + " }}")
    set_cell(t, 5, 1, "{{ wr.lv.bt }}")
    set_cell(t, 5, 2, "{{ wr.lv.at }}")
    wrap_rows(t, 1, 5, "t.winding_res")
    t = T[12]
    for c, k in enumerate(("tap", "pct", "bu", "bv", "bw", "au", "av", "aw")):
        set_cell(t, 3, c, "{{ r." + k + " }}")
    for _ in range(6):
        t._tbl.remove(t.rows[4]._tr)
    tag_row(t, 3, "{%tr for r in ratio_rows %}", "before")
    tag_row(t, 4, "{%tr endfor %}", "after")
    set_cell(T[13], 1, 0, "  Vector group was found to be {{ vector_group }} ")
    t = T[14]
    for c, k in enumerate(("bt.nt", "bt.ht", "bt.lt", "at.nt", "at.ht", "at.lt"), start=1):
        set_cell(t, 3, c, "{{ imp.z." + k + " }}")
        set_cell(t, 4, c, "{{ imp.ll." + k + " }}")
    set_cell(t, 5, 1, "{{ imp.ll50.bt }}")
    set_cell(t, 5, 4, "{{ imp.ll50.at }}")
    for c, name in ((1, "Principal"), (2, "Highest"), (3, "Lowest"), (4, "Principal"), (5, "Highest"), (6, "Lowest")):
        set_cell(t, 2, c, f"{name} ({{{{ tapn.{name} }}}})")
    t = T[15]
    for r, side in ((1, "bt"), (2, "at")):
        for c, k in enumerate(("v", "f", "i", "p"), start=1):
            set_cell(t, r, c, "{{ nl." + side + "." + k + " }}")

    # ======================= Sheet 6: routine contd.
    t = T[17]
    set_cell(t, 1, 1, "{{ eel.tl50.bt }}"); set_cell(t, 1, 2, "{{ eel.tl50.at }}")
    set_cell(t, 2, 1, "{{ eel.tl100.bt }}"); set_cell(t, 2, 2, "{{ eel.tl100.at }}")
    t = T[18]
    for r, k in ((1, "hv_e"), (2, "lv_e"), (3, "hv_lv")):
        set_cell(t, r, 1, "{{ ir.bt." + k + " }}")
        set_cell(t, r, 2, "{{ ir.at." + k + " }}")
    set_cell(T[19], 1, 1, "{{ iov.bt }}"); set_cell(T[19], 1, 2, "{{ iov.at }}")
    set_cell(T[20], 1, 1, "{{ sep.bt }}"); set_cell(T[20], 1, 2, "{{ sep.at }}")

    # ======================= Sheet 7: short circuit
    t = T[22]
    set_cell(t, 10, 1, "{{ sc.condition }}")
    set_cell(t, 17, 1, "{{ circuit_no }}")
    t = T[23]
    for r, tap in ((2, "nt"), (3, "ht"), (4, "lt")):
        set_cell(t, r, 1, "{{ isc." + tap + ".xr }}")
        set_cell(t, r, 2, "{{ isc." + tap + ".rms }}")
        set_cell(t, r, 3, "{{ isc." + tap + ".pk }}")
    t = T[24]
    set_cell(t, 0, 0, "Ability to Withstand the Dynamic Effects of Short-Circuit\n[Sub-clause 4.2 of {{ sc_standard }}]")
    for r in (3, 4, 5, 7, 8, 9, 11, 12, 13, 15):
        cells = unique_cells(t.rows[r])
        sid = re.search(r"S(\d{3})", "".join(p_text(p) for p in cells[0]._tc.iter(qn("w:p")))).group(1)
        key = "osc.s" + sid
        if sid != "011":
            set_cell(t, r, 1, "{{ " + key + ".pk }}")
        set_cell(t, r, 2, "{{ " + key + ".rms }}")
        set_cell(t, r, 3, "{{ " + key + ".dur }}")
        set_cell(t, r, 4, "{{ " + key + ".obs }}")

    # ======================= Sheet 8: reactance, inspection, oil leakage, routine pressure
    t = T[25]
    for r, tap in ((3, "nt"), (4, "ht"), (5, "lt")):
        set_cell(t, r, 1, "{{ xr." + tap + ".b }}")
        set_cell(t, r, 2, "{{ xr." + tap + ".a }}")
        set_cell(t, r, 3, "{{ xr." + tap + ".chg }}")
    set_cell(t, 6, 0, "{{ xr_conclusion }}")
    wrap_rows(t, 1, 7, "t.sc_withstand")
    regex_replace_p(Pe[122], r"(:\s*)$", r"\1{{ sc.insp_conductor }}")
    regex_replace_p(Pe[123], r"(:\s*)$", r"\1{{ sc.insp_spacers }}")
    regex_replace_p(Pe[124], r"(:\s*)$", r"\1 {{ sc.insp_oil }}")
    t = T[26]
    set_cell(t, 2, 1, "{{ ol.atm }} kPa")
    regex_replace_p(unique_cells(t.rows[3])[1]._tc.findall(qn("w:p"))[0], r"^\s+", "{{ ol.amb }} ")
    t = T[27]
    for c, k in enumerate(("head", "test", "hours", "obs")):
        set_cell(t, 1, c, "{{ ol." + k + " }}")
    regex_replace_p(Pe[138], r"(Atmospheric Pressure\s*)kPa", r"\1{{ prr.atm }} kPa")
    regex_replace_p(Pe[139], r"(Ambient temperature)\s+", r"\1      {{ prr.amb }} ")
    t = T[28]
    for c, k in enumerate(("kpa", "min", "obs")):
        set_cell(t, 1, c, "{{ prr." + k + " }}")

    # ======================= Sheet 9: temperature rise
    regex_replace_p(Pe[171], r"\(2623\.20 W\)", "({{ tr.total_loss }} W)")
    regex_replace_p(Pe[172], r"\(14\.58 A\)", "({{ tr.rated_current }} A)")
    regex_replace_p(Pe[177], r"22\.9", "{{ tr.amb }}")
    t = T[30]
    set_cell(t, 0, 1, "{{ tr.nll }}")
    set_cell(t, 1, 1, "{{ tr.fll }}")
    t = T[31]
    for r, k in ((1, "top"), (2, "hv"), (3, "lv")):
        set_cell(t, r, 1, "{{ tr." + k + ".rise }}")
        set_cell(t, r, 2, "{{ tr." + k + ".g }}")
        set_cell(t, r, 3, "{{ tr." + k + ".std }}")
        set_cell(t, r, 4, "{{ tr." + k + ".obs }}")

    # ======================= Sheet 10: type pressure, NLC at 112.5 %
    regex_replace_p(Pe[215], r"complies with", "{{ complies_text }}")
    regex_replace_p(Pe[198], r"91\.0", "{{ prt.atm }}")
    regex_replace_p(Pe[199], r"25\.0", "{{ prt.amb }}")
    for ti, pre in ((33, "prt"), (34, "prv")):
        for c, k in enumerate(("val", "min", "defl", "obs")):
            set_cell(T[ti], 1, c, "{{ " + pre + "." + k + " }}")
    t = T[36]
    set_cell(t, 1, 0, "At 100 percent rated voltage ({{ nlc.v100 }} V)")
    set_cell(t, 1, 1, "At 112.5 percent rated voltage ({{ nlc.v112 }} V)")
    for c, k in enumerate(("lim100", "m100", "lim112", "m112")):
        set_cell(t, 3, c, "{{ nlc." + k + " }}")

    # ======================= generic replacements (body, text boxes, header, footer)
    roots = [doc.element.body] + [s.footer._element for s in doc.sections] + [s.header._element for s in doc.sections]
    for root in roots:
        replace_everywhere(root, r"CPRIBLRSCL25TXXXX(?=[SGP]\d)", "{{ base_no }}")
        replace_everywhere(root, r"CPRIBLRSCL25TXXXX", "{{ report_no }}")
        replace_everywhere(root, r"Date: XXXXXXXXX", "Date: {{ report_date }}")
        replace_everywhere(root, r"X{9}", "{{ engineer_name }}")
        replace_everywhere(root, r"P\. Naveenkumar", "{{ engineer_name }}")
        replace_everywhere(root, r"ULR-TC5452250SCLTXXXX", "{{ ulr_base }}")
        replace_everywhere(root, r"(Principal|Highest|Lowest) \(X\)", r"\1 ({{ tapn.\1 }})")
        replace_everywhere(root, r"\bX \((Principal|Highest|Lowest)\)", r"{{ tapn.\1 }} (\1)")
    # "Refer Sheet 5 of 11" style references - body only (the footer uses PAGE/NUMPAGES fields)
    replace_everywhere(doc.element.body, r"\b(\d{1,2}) of 11\b", r"{{ sh.s\1 }} of {{ n_sheets }}")
    # seal hash in the empty footer cell under "Sheet x of y" (no change in footer height)
    for s in doc.sections:
        ftbl = s.footer.tables[0]
        cell = ftbl.rows[1].cells[2]
        p = cell.paragraphs[0]
        p.alignment = 2
        run = p.add_run("Seal {{ seal12 }}")
        run.font.size = Pt(6.5)
        run.font.name = "Arial"

    # ======================= sheet page breaks + optional blocks
    sheet_start = {2: 7, 3: 14, 4: 25, 5: 54, 6: 74, 7: 104, 8: 115, 9: 149, 10: 191, 11: 226}
    starts = {}
    for sheet, pi in sheet_start.items():
        el = Pe[pi]
        # walk back over padding paragraphs; keep the one closest to "TEST REPORT"
        blanks = []
        prev = el.getprevious()
        while prev is not None and is_blank(prev):
            blanks.append(prev)
            prev = prev.getprevious()
        if blanks:
            keep = blanks[0]
            for b in blanks[1:]:
                b.getparent().remove(b)
        else:
            keep = make_tag_p("")
            el.addprevious(keep)
            keep.remove(keep[0])
        set_page_break_before(keep)
        starts[sheet] = keep
    body_end = doc.element.body.find(qn("w:sectPr"))

    def sheet_last(sheet):
        nxt = starts.get(sheet + 1)
        return nxt.getprevious() if nxt is not None else body_end.getprevious()

    # optional sections inside sheets (wrap before wrapping the whole sheet)
    wrap_block(T[12]._tbl, T[13]._tbl, "t.ratio")
    wrap_block(T[14]._tbl, T[14]._tbl, "t.impedance")
    wrap_block(Pe[64], T[15]._tbl, "t.nll")
    wrap_block(Pe[81], T[17]._tbl, "show.eel")
    wrap_block(Pe[83], T[18]._tbl, "t.ir")
    wrap_block(Pe[85], T[19]._tbl, "t.induced")
    wrap_block(Pe[87], T[20]._tbl, "t.sep_source")
    wrap_block(Pe[118], Pe[125], "t.sc_withstand")
    wrap_block(Pe[126], Pe[127], "show.routine_s8")
    wrap_block(Pe[128], T[27]._tbl, "t.oil_leak")
    wrap_block(Pe[133], T[28]._tbl, "t.pressure_routine")
    wrap_block(Pe[194], T[34]._tbl, "t.pressure_type")
    wrap_block(Pe[205], T[36]._tbl, "t.nlc_112")
    for sheet in (5, 6, 7, 8, 9):
        wrap_block(starts[sheet], sheet_last(sheet), f"show.s{sheet}")

    os.makedirs(os.path.dirname(DST), exist_ok=True)
    doc.save(DST)

    # sanity: nothing left that looks like an unfilled placeholder
    leftovers = []
    for root in roots:
        for p in root.iter(qn("w:p")):
            txt = p_text(p)
            if re.search(r"XXX", txt):
                leftovers.append(txt)
    print("written", DST)
    if leftovers:
        print("WARNING leftover XXX:", leftovers)


if __name__ == "__main__":
    sys.exit(main())
