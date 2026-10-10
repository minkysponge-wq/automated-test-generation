"""
Build sample_data/foreign_layout_losses.xlsx: a losses workbook in a different lab layout - other labels,
shuffled columns, units in the headers, no ATRG placeholders or defined names. It exercises the AI path of
the Excel intake. The expected cell of every value is written to sample_data/foreign_layout_losses.expected.yaml.

    python tools/make_foreign_sample.py
"""
import os

import yaml
from openpyxl import Workbook
from openpyxl.styles import Font

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def main():
    demo = yaml.safe_load(open(os.path.join(ROOT, "sample_data", "demo_values.yaml"), encoding="utf-8"))
    wb = Workbook()
    ws = wb.active
    ws.title = "Loss test record"
    expected = {}

    def put(cell, value, key=None):
        ws[cell] = value
        if key:
            expected[key] = {"cell": f"{ws.title}!{cell}", "value": value}

    ws["A1"] = "Transformer loss test record - Sl. No. 1098, 250 kVA 11000/433 V"
    ws["A1"].font = Font(bold=True)
    # no-load test, columns in another order: Frequency | Loss | Current | Voltage
    ws["A3"] = "Open-circuit (no-load) test at rated voltage"
    for col, h in zip("ABCDE", ["Condition", "f (Hz)", "Loss P0 (W)", "Mean current Io (A)", "Applied voltage (V)"]):
        ws[f"{col}4"] = h
    for r, side, label in ((5, "bt", "Before short-circuit test"), (6, "at", "After short-circuit test")):
        put(f"A{r}", label)
        for col, q in zip("BCDE", ["f", "p", "i", "v"]):
            put(f"{col}{r}", demo[f"nll.{side}.{q}"], f"nll.{side}.{q}")
    # load-loss test per tap: Tap | Stage | W1 W2 W3 | A1 A2 A3 | V1 V2 V3 | Hz
    ws["A9"] = "Short-circuit (load loss) test"
    cols = "ABCDEFGHIJKL"
    for c, h in zip(cols, ["Tap", "Stage", "W1 (W)", "W2 (W)", "W3 (W)", "A1 (A)", "A2 (A)", "A3 (A)",
                           "V1 (V)", "V2 (V)", "V3 (V)", "Hz"]):
        ws[f"{c}10"] = h
    row = 11
    for tap, tlabel in (("nt", "Normal"), ("ht", "Max (+5 %)"), ("lt", "Min (-10 %)")):
        for side, slabel in (("bt", "before SC"), ("at", "after SC")):
            put(f"A{row}", tlabel)
            put(f"B{row}", slabel)
            for c, q in zip(cols[2:], ["p1", "p2", "p3", "i1", "i2", "i3", "v1", "v2", "v3", "f"]):
                put(f"{c}{row}", demo[f"ll.{tap}.{side}.{q}"], f"ll.{tap}.{side}.{q}")
            row += 1
    wb.save(os.path.join(ROOT, "sample_data", "foreign_layout_losses.xlsx"))
    with open(os.path.join(ROOT, "sample_data", "foreign_layout_losses.expected.yaml"), "w", encoding="utf-8") as fh:
        yaml.safe_dump(expected, fh, sort_keys=False)
    print(len(expected), "values written")


if __name__ == "__main__":
    main()
