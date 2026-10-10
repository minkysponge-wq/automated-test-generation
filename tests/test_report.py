"""Generated report (template v1 through the dynamic pipeline): no tags / dummy text left, every value cell filled."""
import re

import docx
import pytest

from app import report
from app.reporting import generator, sections, templates as T

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _all_text(doc):
    parts = [t.text or "" for t in doc.element.body.iter(W + "t")]
    for s in doc.sections:
        parts += [t.text or "" for t in s.footer._element.iter(W + "t")]
        parts += [t.text or "" for t in s.header._element.iter(W + "t")]
    return "".join(parts)


def _ucells(row):
    out, seen = [], set()
    for c in row.cells:
        if id(c._tc) not in seen:
            seen.add(id(c._tc))
            out.append(c)
    return out


# (table, row, unique-cell) that must hold a real value in the full demo report
REQUIRED = ([(0, r, 2) for r in (2, 4, 6, 8, 10, 13, 15, 17, 19, 27, 33, 34, 35, 36, 37, 38)]
            + [(2, r, 2) for r in range(0, 22) if r != 6]
            + [(11, 4, c) for c in range(1, 7)] + [(11, 5, 1), (11, 5, 2)]
            + [(12, r, c) for r in range(3, 10) for c in range(0, 8)]
            + [(14, r, c) for r in (3, 4) for c in range(1, 7)] + [(14, 5, 1), (14, 5, 4)]
            + [(15, r, c) for r in (1, 2) for c in range(1, 5)]
            + [(17, r, c) for r in (1, 2) for c in (1, 2)]
            + [(18, r, c) for r in (1, 2, 3) for c in (1, 2)]
            + [(19, 1, 1), (19, 1, 2), (20, 1, 1), (20, 1, 2)]
            + [(23, r, c) for r in (2, 3, 4) for c in (1, 2, 3)]
            + [(24, r, c) for r in (3, 4, 5, 7, 8, 9, 11, 12, 13) for c in (1, 2, 3)]
            + [(25, r, c) for r in (3, 4, 5) for c in (1, 2, 3)]
            + [(27, 1, c) for c in range(4)] + [(28, 1, c) for c in range(3)]
            + [(30, 0, 1), (30, 1, 1)] + [(31, r, c) for r in (1, 2, 3) for c in range(1, 5)]
            + [(33, 1, c) for c in range(4)] + [(34, 1, c) for c in range(4)]
            + [(36, 3, c) for c in range(4)])


@pytest.fixture(scope="module")
def demo_docx(seeded, tmp_path_factory):
    conn = seeded["conn"]
    snap = report.build_snapshot(conn, seeded["review_job_id"])
    out = tmp_path_factory.mktemp("rep") / "demo.docx"
    report.render_docx(conn, snap, str(out), {"report_date": "18-11-2025", "reviewer_name": "HOD", "seal": "a" * 64})
    return docx.Document(str(out))


def test_no_placeholders_left(demo_docx):
    text = _all_text(demo_docx)
    assert "{{" not in text and "{%" not in text and "}}" not in text
    assert "XXXX" not in text
    assert "CPRIBLRSCL25T1655" in text and "HVD25S0848" in text


def test_required_cells_filled(demo_docx):
    t = demo_docx.tables
    assert len(t) == 39
    empty = []
    for ti, r, c in REQUIRED:
        val = _ucells(t[ti].rows[r])[c].text.strip()
        if not val or re.fullmatch(r"[—\-\s]+", val):
            empty.append((ti, r, c, val))
    assert not empty, empty


def test_key_values_match_logsheets(demo_docx):
    t = demo_docx.tables
    assert _ucells(t[31].rows[2])[1].text.strip() == "39.7"        # HV winding rise
    assert _ucells(t[31].rows[3])[1].text.strip() == "34.3"        # LV winding rise
    assert _ucells(t[25].rows[3])[3].text.strip() == "-0.163"      # % change in X principal tap
    assert _ucells(t[36].rows[3])[3].text.strip() == "1.28 %"      # NLC at 112.5 %
    assert _ucells(t[17].rows[1])[1].text.strip() == "908.09"      # total loss 50 %
    assert "Within limits" in _ucells(t[31].rows[1])[4].text


def test_dry_run_of_active_template_is_clean(seeded, tmp_path):
    rec = T.active(seeded["conn"])
    out = tmp_path / "dry.docx"
    generator.render(T._abs(rec["compiled_path"]), T.mapping_of(rec), T.demo_snapshot(), str(out))
    assert generator.unfilled(str(out)) == []


def test_sheet_numbers_follow_requested_tests(seeded):
    meta = T.mapping_of(T.active(seeded["conn"]))["compiled_meta"]
    show, flags, numbers, n, test_sheets = sections.visibility(meta, ["winding_res", "oil_leak"])
    assert n == 8                                   # cover..drawings, sheet 5, sheet 8, sheet 10, notes
    assert not flags[7] and not flags[9] and flags[10]
    assert numbers[8] == 6 and numbers[11] == 8
    assert test_sheets == {"winding_res": [5], "oil_leak": [6]}
