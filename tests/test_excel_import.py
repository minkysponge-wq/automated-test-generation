"""Excel template -> fill -> import round trip, driven by mappings/fields.yaml."""
import os

import yaml

import config
from app import excel_io, fields as F

ALL_TESTS = list(F.tests())


def _demo():
    with open(os.path.join(config.SAMPLE_DATA_DIR, "demo_values.yaml"), encoding="utf-8") as fh:
        return yaml.safe_load(fh)


def test_round_trip_every_station(tmp_path):
    demo = _demo()
    for st, meta in F.stations().items():
        wb = excel_io.build_station_template(st)
        vals = {k: v for k, v in demo.items() if F.field(k) and F.field(k)["station"] == st}
        excel_io.fill_values(wb, vals)
        path = tmp_path / meta["file"]
        wb.save(path)
        res = excel_io.import_workbook(str(path), st, ALL_TESTS)
        assert res.ok, res.errors
        assert not res.missing and not res.unfilled and not res.non_numeric, (st, res.summary())
        for k, v in vals.items():
            got = res.values[k][0]
            if F.field(k)["type"] == "number":
                assert float(got) == float(v), (k, got, v)
            else:
                assert got == str(v), (k, got, v)
            assert "!" in res.values[k][1]           # source cell reference recorded


def test_blank_template_reports_unfilled_placeholders():
    path = os.path.join(config.EXCEL_TEMPLATE_DIR, "losses.xlsx")
    res = excel_io.import_workbook(path, "losses", ALL_TESTS)
    assert not res.ok
    assert len(res.unfilled) > 50


def test_falls_back_to_token_scan_when_names_are_lost(tmp_path):
    wb = excel_io.build_station_template("pressure")
    excel_io.fill_values(wb, {"prt.kpa": 80})
    for name in list(wb.defined_names.keys()):
        if name != "prt_kpa":
            del wb.defined_names[name]
    p = tmp_path / "noname.xlsx"
    wb.save(p)
    res = excel_io.import_workbook(str(p), "pressure", ["pressure_type"])
    assert res.values["prt.kpa"][0] == "80"
    unfilled_keys = {k for k, _ in res.unfilled}
    assert "prt.min" in unfilled_keys          # found by its leftover {{token}}


def test_layout_change_does_not_break_import(tmp_path):
    """Move a filled cell somewhere else (and its defined name with it, as Excel does) - still imported."""
    from openpyxl.workbook.defined_name import DefinedName
    wb = excel_io.build_station_template("pressure")
    excel_io.fill_values(wb, {"ol.obs": "No leakage at any point"})
    ws = wb["Logsheet"]
    dn = wb.defined_names["ol_obs"]
    title, coord = next(iter(dn.destinations))
    ws["Q40"] = ws[coord.replace("$", "")].value
    ws[coord.replace("$", "")] = None
    wb.defined_names["ol_obs"] = DefinedName("ol_obs", attr_text="'Logsheet'!$Q$40")
    p = tmp_path / "moved.xlsx"
    wb.save(p)
    res = excel_io.import_workbook(str(p), "pressure", ["oil_leak"])
    assert res.values["ol.obs"] == ("No leakage at any point", "Logsheet!Q40")


def test_non_numeric_and_wrong_sample_code(tmp_path):
    wb = excel_io.build_station_template("pressure")
    excel_io.fill_values(wb, {"prt.kpa": "eighty", "pressure.hdr.sample_code": "HVD25S0001"})
    p = tmp_path / "bad.xlsx"
    wb.save(p)
    res = excel_io.import_workbook(str(p), "pressure", ["pressure_type"], expected_sample_code="HVD25S0847")
    assert ("prt.kpa", res.values["prt.kpa"][1], "eighty") in res.non_numeric
    assert res.errors and "HVD25S0001" in res.errors[0]


def test_sample_data_imports_cleanly():
    for st, meta in F.stations().items():
        res = excel_io.import_workbook(os.path.join(config.SAMPLE_DATA_DIR, meta["file"]), st, ALL_TESTS, "HVD26S0847")
        assert res.ok and not res.unfilled and not res.missing and not res.non_numeric, (st, res.summary())


def test_request_excel_round_trip():
    from app import request_form
    vals = excel_io.parse_request_workbook(os.path.join(config.SAMPLE_DATA_DIR, "customer_request_filled.xlsx"))
    out = request_form.from_excel(vals, ALL_TESTS)
    assert out["fields"]["serial_numbers"] == "1098"
    assert len(out["drawings"]) == 2 and set(out["tests"]) == set(ALL_TESTS)
