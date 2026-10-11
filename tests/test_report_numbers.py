"""
Sample code and test report number are allocated automatically and reach the CPRI template fields.

In this version both numbers are allocated when the Test In-charge accepts the request
(jobs.accept); nobody types either one. The issued demo report (sealed by the existing HOD
issue function during seeding) is opened with python-docx and checked field by field.
"""
import json
import os
import re

import docx

import config
from app import jobs
from app.db import tx
from conftest import user

SAMPLE_RE = re.compile(r"^HVD\d{2}S\d{4}$")
REPORT_RE = re.compile(r"^CPRIBLRSCL\d{2}T\d{4}$")
TEMPLATE = os.path.join(config.BASE_DIR, "reference", "Transformer_Test_report_format_1.docx")


def _cover_value(doc, label):
    for row in doc.tables[0].rows:
        cells = [c.text.strip() for c in row.cells]
        if cells and cells[0].startswith(label):
            return cells[-1]
    raise AssertionError(f"cover row '{label}' not found")


def _sheet_headers(doc):
    return [c.text.strip() for t in doc.tables for row in t.rows for c in row.cells
            if c.text.strip().startswith("Test Report Number:")]


def test_numbers_allocated_without_typing(conn):
    cust, admin = user(conn, "customer1"), user(conn, "admin")
    with tx(conn):
        job_id, ref = jobs.create_request(conn, cust, {"customer_name": "Number check", "rating": "100 kVA"}, ["T1"])
    job = jobs.get_job(conn, job_id)
    assert ref.startswith("REQ-") and job["sample_code"] is None and job["report_no"] is None
    with tx(conn):
        sample_code, report_no = jobs.accept(conn, job_id, admin, {"physical_condition": "Suitable",
                                                                   "lab_capability": "Yes"}, {}, job["version"])
    assert SAMPLE_RE.match(sample_code) and REPORT_RE.match(report_no)
    job = jobs.get_job(conn, job_id)
    assert (job["sample_code"], job["report_no"]) == (sample_code, report_no)
    assert json.loads(job["wi_json"])["sample_code"] == sample_code


def test_issued_report_carries_both_numbers(seeded, conn):
    job = jobs.get_job(conn, seeded["job_id"])
    rep = conn.execute("SELECT * FROM reports WHERE job_id=?", (job["id"],)).fetchone()
    assert rep is not None and REPORT_RE.match(rep["report_no"]) and SAMPLE_RE.match(job["sample_code"])
    doc = docx.Document(rep["docx_path"])

    cover_no = _cover_value(doc, "Test Report Number")
    assert cover_no.split()[0] == rep["report_no"] and "XXXX" not in cover_no
    cover_sample = _cover_value(doc, "CPRI Sample Code Number")
    assert cover_sample == job["sample_code"]

    headers = _sheet_headers(doc)
    expected = len(_sheet_headers(docx.Document(TEMPLATE)))     # every per-sheet header the template has
    assert expected and len(headers) == expected
    for h in headers:
        assert f"Test Report Number: {rep['report_no']}" in h and "XXXX" not in h


def test_download_name_has_both_numbers(seeded, conn, client):
    from conftest import login
    job = jobs.get_job(conn, seeded["job_id"])
    owner = conn.execute("SELECT username FROM users WHERE id=?", (job["customer_id"],)).fetchone()["username"]
    assert owner == "customer1"
    login(client, "customer", "customer1", "cust12345")
    r = client.get(f"/customer/requests/{job['id']}/report.docx")
    assert r.status_code == 200
    assert f"{job['report_no']}_{job['sample_code']}.docx" in r.headers["Content-Disposition"]
