"""
Compiler: source DOCX + mapping -> docxtpl-tagged DOCX (all original formatting kept).

Generated tags (filled by generator.py):
    {{ F['key'] }}                 registry field / computed / system value
    {{ X['slot'] }}                expression slot (evaluated safely in Python, not in Jinja)
    {%tr for it in L['list'] %}    repeating table row, cells {{ it.attr }}
    {%tr if T['code'] %}           schedule-of-tests row printed only for requested tests
    {%tr if G['group'] %}          group heading rows of the schedule
    {{ SCHED['code'] }}            "Sheet" column of the schedule
    {%p if SHOW['id'] %} / {%tr if SHOW['id'] %}   sections
    {%p if SHEET[n] %}             whole sheets
    {{ SH[n] }} of {{ NSHEETS }}   sheet references ("Refer Sheet 5 of 11")
"""
import re

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

from . import sections as S
from .docx_parser import ParsedDoc
from .docxutil import (is_blank_p, make_tag_p, p_text, regex_replace_p, set_page_break_before, set_tc_text,
                       tag_row, unique_cells)

FIELD_TOKEN = re.compile(r"\{\{F:([A-Za-z0-9_.]+)\}\}")


def _tag(binding, slot_id):
    kind = binding.get("kind")
    if kind == "field":
        return "{{ F['%s'] }}" % binding["key"]
    if kind == "expr":
        return "{{ X['%s'] }}" % slot_id
    if kind == "text":
        return (binding.get("text") or "").replace("{", "(").replace("}", ")")
    return None


def _repl_template(text):
    """Manual-rule replacement text: {{F:key}} -> Jinja tag; other braces neutralised."""
    out = FIELD_TOKEN.sub(lambda m: "\x00%s\x01" % m.group(1), text)
    out = out.replace("{", "(").replace("}", ")")
    return out.replace("\x00", "{{ F['").replace("\x01", "'] }}")


def compile_template(source_path, mapping, out_path):
    doc = Document(source_path)
    pd = ParsedDoc(doc, mapping.get("sheet_marker", r"^TEST REPORT$"))
    parsed = pd.to_json()
    warnings = []
    slots_by_id = {s["id"]: s for s in parsed["slots"]}
    bindings = mapping.get("slots", {})

    # resolve text anchors BEFORE editing (element references stay valid afterwards)
    resolved, w = S.resolve(parsed, mapping.get("sections", []))
    warnings += w
    sheet_start_els = {sh["n"]: pd.block_refs[sh["start_block"]] for sh in parsed["sheets"]}
    block_sheet = {b["i"]: b["sheet"] for b in parsed["blocks"]}

    # ---- 1. slots
    per_el = {}
    for sid, b in bindings.items():
        s = slots_by_id.get(sid)
        if s is None:
            warnings.append(f"Mapped slot {sid} no longer exists in this document")
            continue
        tag = _tag(b, sid)
        if tag is None:
            continue
        per_el.setdefault(id(pd.refs[sid]), []).append((s, tag))
    compiled = 0
    for items in per_el.values():
        items.sort(key=lambda x: -x[0].get("occ", 0))
        for s, tag in items:
            el = pd.refs[s["id"]]
            k = s["kind"]
            if k in ("cell_empty", "cell_dummy", "obs_cell"):
                set_tc_text(el, tag)
            elif k == "unit_cell":
                p0 = el.findall(qn("w:p"))[0]
                if not regex_replace_p(p0, r"^\s*", tag + " ", count=1):
                    set_tc_text(el, tag)
            elif k in ("span", "placeholder", "inline_number"):
                regex_replace_p(el, re.escape(s["find"]), lambda m, t=tag: t, only_occurrence=s.get("occ", 0))
            elif k == "after_colon":
                regex_replace_p(el, r"(:\s*)$", lambda m, t=tag: m.group(1) + ("" if m.group(1).endswith(" ") else " ") + t,
                                count=1)
            compiled += 1

    # ---- 2. manual rules
    body = doc.element.body
    tables = pd.table_objs
    for rule in mapping.get("replacements", []):
        if rule.get("kind") == "cell":
            bi = S.find_block(parsed["blocks"], {"type": "t", "text": rule["table"]})
            if bi is None:
                warnings.append(f"Rule cell: table '{rule['table']}' not found")
                continue
            t = next(x for x in parsed["tables"] if x["block"] == parsed["blocks"][bi]["i"])
            try:
                cell = unique_cells(tables[t["t"]].rows[rule["row"]])[rule["col"]][1]
            except IndexError:
                warnings.append(f"Rule cell: cell {rule['row']},{rule['col']} not found")
                continue
            set_tc_text(cell._tc, "{{ F['%s'] }}" % rule["field"])
            continue
        roots = [body]
        scope = rule.get("scope") or {}
        if scope.get("block"):
            bi = S.find_block(parsed["blocks"], {"text": scope["block"], "type": scope.get("type")})
            if bi is None:
                warnings.append(f"Rule '{rule['pattern']}': scope '{scope['block']}' not found")
                continue
            roots = [pd.block_refs[bi]]
            if scope.get("row_text") and parsed["blocks"][bi]["type"] == "t":
                t = next(x for x in parsed["tables"] if x["block"] == parsed["blocks"][bi]["i"])
                roots = [r._tr for r in tables[t["t"]].rows if scope["row_text"].lower() in
                         " ".join(c.text for _, c in unique_cells(r)).lower()]
        repl = _repl_template(rule["replace"])
        n = 0
        for root in roots:
            for p in list(root.iter(qn("w:p"))):
                n += regex_replace_p(p, rule["pattern"], repl)
        if not n and not rule.get("optional"):
            warnings.append(f"Rule '{rule['pattern']}' matched nothing")

    # ---- 3. repeating rows
    for lp in mapping.get("loops", []):
        if lp["t"] >= len(tables):
            warnings.append(f"Loop {lp['id']}: table missing")
            continue
        tbl = tables[lp["t"]]
        rows = list(tbl.rows)
        tpl_row = rows[lp["first_row"]]
        for (gi, c), attr in zip(unique_cells(tpl_row), lp["columns"]):
            set_tc_text(c._tc, "{{ it.%s }}" % attr if attr else "")
        for r in rows[lp["first_row"] + 1:lp["last_row"] + 1]:
            tbl._tbl.remove(r._tr)
        tag_row(tpl_row._tr, "{%%tr for it in L['%s'] %%}" % lp["list"], "before")
        tag_row(tpl_row._tr, "{%tr endfor %}", "after")

    # ---- 4. schedule of tests
    schedule_tests = []
    for sch in mapping.get("schedule", []):
        tbl = tables[sch["t"]]
        rows = list(tbl.rows)
        for r_str, info in sorted(sch["rows"].items(), key=lambda kv: -int(kv[0])):
            tr = rows[int(r_str)]
            if info["kind"] == "test":
                schedule_tests.append(info["code"])
                if sch.get("sheet_col") is not None:
                    for gi, c in unique_cells(tr):
                        if gi == sch["sheet_col"]:
                            set_tc_text(c._tc, "{{ SCHED['%s'] }}" % info["code"])
                cond = "T['%s']" % info["code"]
            else:
                cond = "G['%s']" % info["group"]
                if info["group"] == "routine":
                    c = unique_cells(tr)[0][1]
                    if "routine" in c.text.lower():
                        set_tc_text(c._tc, "{{ F['sys.routine_heading'] }}")
            tag_row(tr._tr, "{%tr endif %}", "after")
            tag_row(tr._tr, "{%%tr if %s %%}" % cond, "before")

    # ---- 5. "Sheet N of M" references in the body
    m_total = len(parsed["sheets"])
    if m_total:
        pat = r"\b(\d{1,2}) of %d\b" % m_total
        for p in list(body.iter(qn("w:p"))):
            regex_replace_p(p, pat, lambda m: "{{ SH[%d] }} of {{ NSHEETS }}" % int(m.group(1)))

    # ---- 6. footer seal (empty cell under the "Sheet x of y" cell)
    if mapping.get("footer_seal", True):
        for sec in doc.sections:
            placed = False
            for t in sec.footer.tables:
                for ri, row in enumerate(t.rows):
                    for gi, cell in enumerate(row.cells):
                        if cell._tc.find(".//" + qn("w:instrText")) is not None and "PAGE" in "".join(
                                x.text or "" for x in cell._tc.iter(qn("w:instrText"))):
                            if ri + 1 < len(t.rows):
                                below = t.rows[ri + 1].cells[gi]
                                if not below.text.strip():
                                    p = below.paragraphs[0]
                                    p.alignment = 2
                                    run = p.add_run("Seal {{ F['sys.seal12'] }}")
                                    run.font.size = Pt(6.5)
                                    run.font.name = "Arial"
                                    placed = True
                        if placed:
                            break
                    if placed:
                        break
            if not placed:
                warnings.append("No empty footer cell found for the seal hash")

    # ---- 7. one sheet per page: explicit page break, padding removed
    sheet_heads = {}
    for n, el in sheet_start_els.items():
        if n == 1:
            sheet_heads[n] = el
            continue
        blanks, prev = [], el.getprevious()
        while prev is not None and is_blank_p(prev):
            blanks.append(prev)
            prev = prev.getprevious()
        if blanks:
            keep = blanks[0]
            for b in blanks[1:]:
                b.getparent().remove(b)
        else:
            keep = make_tag_p("")
            keep.remove(keep[0])
            el.addprevious(keep)
        set_page_break_before(keep)
        sheet_heads[n] = keep

    # ---- 8. sections (rows first, then blocks), then whole sheets
    for sec in [x for x in resolved if x["kind"] == "rows"]:
        rows = list(tables[sec["t"]].rows)
        tag_row(rows[sec["to_row"]]._tr, "{%tr endif %}", "after")
        tag_row(rows[sec["from_row"]]._tr, "{%%tr if SHOW['%s'] %%}" % sec["id"], "before")
    for sec in [x for x in resolved if x["kind"] == "block"]:
        first, last = pd.block_refs[sec["start"]], pd.block_refs[sec["end"]]
        first.addprevious(make_tag_p("{%%p if SHOW['%s'] %%}" % sec["id"]))
        last.addnext(make_tag_p("{%p endif %}"))
    keep_texts = [t.lower() for t in mapping.get("keep_sheets_with", [])]
    sheets_meta = []
    for sh in parsed["sheets"]:
        n = sh["n"]
        secs = [x["id"] for x in resolved if x["sheet"] == n]
        text = " ".join(b["text"] for b in parsed["blocks"] if b["sheet"] == n).lower()
        keep = any(k in text for k in keep_texts)
        sheets_meta.append({"n": n, "sections": secs, "keep": keep})
        if secs and not keep and n in sheet_heads:
            sheet_heads[n].addprevious(make_tag_p("{%%p if SHEET[%d] %%}" % n))
            nxt = sheet_heads.get(n + 1)
            end_el = nxt.getprevious() if nxt is not None else body.find(qn("w:sectPr")).getprevious()
            end_el.addnext(make_tag_p("{%p endif %}"))

    doc.save(out_path)
    meta = {"sections": [{"id": x["id"], "tests": x["tests"], "sheet": x["sheet"], "schedule": x.get("schedule", True)}
                         for x in resolved],
            "sheets": sheets_meta, "source_sheet_count": m_total, "loops": [l["list"] for l in mapping.get("loops", [])],
            "schedule_tests": schedule_tests, "compiled_slots": compiled, "warnings": warnings,
            "expr_slots": {sid: b["expr"] for sid, b in bindings.items() if b.get("kind") == "expr"}}
    return meta
