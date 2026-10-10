"""
Parse the structure of any report DOCX.

Output (JSON-serialisable):
    blocks  - body paragraphs and tables in order, with sheet number and a text fingerprint
    sheets  - sheet boundaries (a sheet starts at a paragraph matching the sheet marker, e.g. "TEST REPORT")
    slots   - places where a value goes:
        cell_empty    empty value cell in a table (required when the row has a label)
        cell_dummy    cell made of dummy text (XXXX..., several tokens) -> whole cell replaced
        span          one dummy token inside text ("Test Report Number: CPRIBLRSCL25TXXXX")
        placeholder   an existing {{field.key}} token
        unit_cell     cell holding only a unit ("  kPa") -> value goes in front of the unit
        obs_cell      cell under an "Observation(s)" column (default text kept unless mapped)
        inline_number decimal number inside a body paragraph ("Total loss (2623.20 W)")
        after_colon   body paragraph "Label :" with nothing after the colon
      each slot has a context {row, col, title, heading, para} used by the auto-mapper and a
      fingerprint used to carry mappings from one template version to the next.
    tables  - per table: fingerprint, loop candidate (rows numbered 1..N), schedule-of-tests detection
"""
import re

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table

from .docxutil import all_text, norm, p_text, unique_cells

DUMMY_RE = re.compile(r"[\w\-/]*?X{3,}")
PLACEHOLDER_RE = re.compile(r"\{\{\s*([A-Za-z0-9_.]+)\s*\}\}")
UNIT_ONLY_RE = re.compile(r"^\s*(kPa|°C|C|V|A|W|Hz|mm|%|kV|mΩ|Ω|kA|s|h)\s*$")
INLINE_NUM_RE = re.compile(r"(?<![\w.\-/])\d+\.\d+(?![\w.\-/])")
AFTER_COLON_RE = re.compile(r"^(?P<label>[^:]{3,60}?)\s*:\s*$")
CLAUSE_RE = re.compile(r"\d+\.\d+\s*[a-z]\)")


def _is_heading(text):
    t = text.strip()
    if DUMMY_RE.search(t):
        return False
    letters = [c for c in t if c.isalpha()]
    return 4 <= len(t) <= 70 and letters and sum(c.isupper() for c in letters) / len(letters) > 0.85


def _clean_label(text):
    if re.fullmatch(r"[\s\-–—]*", text or ""):
        return ""
    t = DUMMY_RE.sub(" ", text or "")
    t = PLACEHOLDER_RE.sub(" ", t)
    return re.sub(r"\s+", " ", t).strip(" :\t")


class ParsedDoc:
    """Holds the JSON structure plus live element references (for the compiler)."""

    def __init__(self, doc, sheet_marker=r"^TEST REPORT$"):
        self.doc = doc
        self.marker = re.compile(sheet_marker, re.I)
        self.blocks, self.slots, self.tables, self.sheets = [], [], [], []
        self.refs = {}          # slot id -> lxml element (w:tc or w:p)
        self.block_refs = []    # block index -> element
        self.table_objs = []    # table index -> docx Table
        self._parse()

    # ------------------------------------------------------------------ helpers
    def _add_slot(self, el, **s):
        s.setdefault("required", True)
        s.setdefault("find", None)
        s.setdefault("occ", 0)
        ctx = s.setdefault("context", {})
        parts = [s["kind"] if s["kind"] in ("inline_number", "after_colon", "unit_cell") else "v",
                 norm(ctx.get("title", ""))[:60], norm(ctx.get("col", ""))[:80], norm(ctx.get("row", ""))[:80],
                 norm(ctx.get("para", ""))[:80]]
        base = "|".join(parts)
        n = sum(1 for x in self.slots if x["fp"].rsplit("#", 1)[0] == base)
        s["fp"] = f"{base}#{n}"
        self.slots.append(s)
        self.refs[s["id"]] = el

    def _spans(self, p, text, base_id, ctx, where, fallback_label=""):
        """Dummy tokens and placeholders inside one paragraph."""
        k = 0
        counts = {}
        for m in DUMMY_RE.finditer(text):
            tok = m.group(0)
            occ = counts.get(tok, 0)
            counts[tok] = occ + 1
            para_ctx = dict(ctx, para=text.replace(tok, "#"))
            before = _clean_label(text[:m.start()])
            after_raw = DUMMY_RE.sub(" ", text[m.end():]).strip()
            after = "" if after_raw.endswith(":") or ":" in after_raw[:-1] else _clean_label(after_raw)
            # nearest label wins: text before the token, then after it, then the rest of the cell, then the row
            label = before or after or _clean_label(fallback_label) or ctx.get("row", "")
            if label:
                para_ctx["row"] = label
            self._add_slot(p, id=f"{base_id}s{k}", kind="span", find=tok, occ=occ, current=tok, context=para_ctx,
                           **self._loc)
            k += 1
        for m in PLACEHOLDER_RE.finditer(text):
            self._add_slot(p, id=f"{base_id}h{k}", kind="placeholder", find=m.group(0), occ=0, current=m.group(0),
                           placeholder=m.group(1), context=dict(ctx, para=text), **self._loc)
            k += 1

    # ------------------------------------------------------------------ parse
    def _parse(self):
        body = self.doc.element.body
        sheet, heading = 0, []
        last_para_text = ""
        ti = 0
        children = list(body.iterchildren())
        for bi, el in enumerate(children):
            tag = el.tag.split("}")[1]
            if tag == "p":
                text = p_text(el)
                if self.marker.match(text.strip()):
                    sheet += 1
                    self.sheets.append({"n": sheet, "start_block": bi})
                    heading = []
                elif _is_heading(text):
                    heading = (heading + [text.strip()])[-2:]
                self.blocks.append({"i": bi, "type": "p", "text": text[:200], "fp": norm(text)[:120],
                                    "sheet": max(sheet, 1)})
                self.block_refs.append(el)
                if text.strip():
                    last_para_text = text.strip()
                self._loc = {"part": "body", "block": bi, "sheet": max(sheet, 1)}
                ctx = {"heading": " ".join(heading)}
                nxt = ""
                for el2 in children[bi + 1:bi + 4]:
                    if el2.tag == qn("w:p") and p_text(el2).strip():
                        nxt = p_text(el2)
                        break
                self._body_para_slots(el, text, bi, ctx, nxt)
            elif tag == "tbl":
                table = Table(el, self.doc)
                self.table_objs.append(table)
                info = self._table(table, ti, bi, max(sheet, 1), last_para_text, " ".join(heading))
                self.tables.append(info)
                self.blocks.append({"i": bi, "type": "t", "t": ti, "text": info["text"][:200], "fp": info["fp"],
                                    "sheet": max(sheet, 1)})
                self.block_refs.append(el)
                # a table's first row can be a heading too ("SPECIAL TEST")
                first = info["rows"][0][0] if info["rows"] and info["rows"][0] else ""
                if _is_heading(first):
                    heading = (heading + [first.strip()])[-2:]
                ti += 1
            else:
                self.blocks.append({"i": bi, "type": tag, "text": "", "fp": "", "sheet": max(sheet, 1)})
                self.block_refs.append(el)
        for i, s in enumerate(self.sheets):
            s["end_block"] = self.sheets[i + 1]["start_block"] - 1 if i + 1 < len(self.sheets) else len(self.blocks) - 1
        self._footer_slots()

    def _body_para_slots(self, el, text, bi, ctx, next_text=""):
        if not text.strip():
            return
        self._spans(el, text, f"b{bi}", ctx, "para", next_text)
        if not DUMMY_RE.search(text):
            # detect on text with tabs as spaces ("Pressure<tab>91.0 kPa"), locate on the raw run text
            tabbed = tab_text(el)
            seen = {}
            for k, m in enumerate(INLINE_NUM_RE.finditer(tabbed)):
                num = m.group(0)
                occ = seen.get(num, 0)
                seen[num] = occ + 1
                self._add_slot(el, id=f"b{bi}n{k}", kind="inline_number", find=num, occ=occ, current=num,
                               required=False, context=dict(ctx, row=tabbed.replace(num, " "), para=tabbed),
                               **self._loc)
        m = AFTER_COLON_RE.match(text.replace("\t", " "))
        if m and not DUMMY_RE.search(text):
            self._add_slot(el, id=f"b{bi}c", kind="after_colon", current="", required=False,
                           context=dict(ctx, row=m.group("label"), para=text), **self._loc)

    def _table(self, table, ti, bi, sheet, prev_para, heading):
        rows_text = []
        grid = []           # grid[r][gi] = (tc, text)
        for row in table.rows:
            cells = row.cells
            grid.append([(c._tc, c.text.strip()) for c in cells])
            rows_text.append([t for _, t in unique_cells_text(row)])
        title = ""
        if rows_text and len(set(rows_text[0])) == 1 and rows_text[0][0]:
            title = rows_text[0][0]
        elif prev_para:
            title = prev_para
        text = " ".join(" ".join(r) for r in rows_text)
        fp = norm(" ".join(rows_text[0]) if rows_text else "")[:120]
        cells = [[{"gi": gi, "text": c.text.strip()} for gi, c in unique_cells(row)] for row in table.rows]
        info = {"t": ti, "block": bi, "fp": fp, "text": text, "rows": rows_text, "cells": cells, "n_rows": len(rows_text),
                "loop": None, "schedule": None}
        self._detect_loop(table, info, grid)
        self._detect_schedule(info)
        loop_rows = set(range(info["loop"]["first_row"], info["loop"]["last_row"] + 1)) if info["loop"] else set()
        sched_sheet_col = info["schedule"]["sheet_col"] if info["schedule"] else None
        for r, row in enumerate(table.rows):
            if r in loop_rows:
                continue
            ucells = unique_cells(row)
            row_texts = [c.text.strip() for _, c in ucells]
            data_row = False
            if not any(row_texts):
                # an all-empty row right under a header row (several headers, no ":" label column) is a data row
                above = [c.text.strip() for _, c in unique_cells(table.rows[r - 1])] if r > 0 else []
                data_row = (r == len(table.rows) - 1 and len([t for t in above if t]) >= 2 and ":" not in above
                            and len(ucells) >= 2)
                if not data_row:
                    continue
            elif all(t in ("", ":") for t in row_texts):
                continue
            signature_row = any(DUMMY_RE.search(t) for t in row_texts) and not any(
                t and not DUMMY_RE.search(t) and t != ":" for t in row_texts[1:])
            for ci, (gi, cell) in enumerate(ucells):
                text = cell.text.strip()
                if sched_sheet_col is not None and r > 0 and gi == sched_sheet_col:
                    continue
                row_label = ""
                for gj, c2 in reversed(ucells[:ci]):
                    t2 = c2.text.strip()
                    if t2 and t2 != ":" and not DUMMY_RE.fullmatch(t2) and not re.fullmatch(r"[\s\-–—]+", t2):
                        row_label = _clean_label(t2)
                        break
                col_parts = []
                for rr in range(r - 1, -1, -1):
                    if gi >= len(grid[rr]):
                        continue
                    tc2, t2 = grid[rr][gi]
                    if tc2 is cell._tc or not t2 or DUMMY_RE.search(t2) or (text and t2 == text):
                        continue
                    if len({id(x[0]) for x in grid[rr]}) == 1 and len(grid[rr]) > 1:
                        continue        # merged across the full width: a group/title row, not a column header
                    if t2 not in col_parts:
                        col_parts.append(t2)
                    if len(col_parts) >= 3:
                        break
                col = " / ".join(reversed(col_parts))
                ctx = {"row": row_label, "col": col, "title": title, "heading": heading}
                self._loc = {"part": "body", "block": bi, "sheet": sheet, "t": ti, "row": r, "col": gi}
                sid = f"t{ti}r{r}c{gi}"
                tokens = DUMMY_RE.findall(text)
                residual = DUMMY_RE.sub("", text)
                residual = re.sub(r"[\s,/&|:.\-]+", "", residual)
                if PLACEHOLDER_RE.search(text):
                    for pi, p in enumerate(cell._tc.findall(qn("w:p"))):
                        self._spans(p, p_text(p), f"{sid}p{pi}", ctx, "cell")
                elif tokens and (not residual or (len(tokens) >= 2 and ":" not in DUMMY_RE.sub("", text))):
                    self._add_slot(cell._tc, id=sid, kind="cell_dummy", current=text, context=ctx, **self._loc)
                elif tokens:
                    paras = cell._tc.findall(qn("w:p"))
                    for pi, p in enumerate(paras):
                        others = " ".join(p_text(x) for x in paras if x is not p)
                        self._spans(p, p_text(p), f"{sid}p{pi}", ctx, "cell", others)
                elif UNIT_ONLY_RE.match(cell.text or "") and gi > 0 and len(row_label) >= 3:
                    self._add_slot(cell._tc, id=sid, kind="unit_cell", current=text, context=ctx, **self._loc)
                elif text == "" and (gi > 0 or data_row) and not signature_row:
                    self._add_slot(cell._tc, id=sid, kind="cell_empty", current="", context=ctx,
                                   required=bool(row_label or col), **self._loc)
                elif text and r > 0 and re.search(r"observation", col, re.I) and not re.search(r"observation", text, re.I):
                    self._add_slot(cell._tc, id=sid, kind="obs_cell", current=text, context=ctx, required=False,
                                   **self._loc)
        return info

    def _detect_loop(self, table, info, grid):
        """Rows whose first cell is 1, 2, 3 ... and whose other cells are empty -> repeating row."""
        rows = info["rows"]
        run = []
        for r, cells in enumerate(rows):
            if cells and cells[0].strip().isdigit() and all(not c.strip() for c in cells[1:]):
                if not run or int(cells[0]) == int(rows[run[-1]][0]) + 1:
                    run.append(r)
                    continue
            if len(run) >= 3:
                break
            run = []
        if len(run) >= 3:
            first = run[0]
            headers = []
            for gi in range(len(grid[first])):
                parts = []
                for rr in range(first - 1, -1, -1):
                    if gi < len(grid[rr]) and grid[rr][gi][1] and grid[rr][gi][1] not in parts:
                        parts.append(grid[rr][gi][1])
                    if len(parts) >= 2:
                        break
                headers.append(" ".join(reversed(parts)))
            uniq = []
            seen = set()
            for gi, (tc, _t) in enumerate(grid[first]):
                if id(tc) in seen:
                    continue
                seen.add(id(tc))
                uniq.append(headers[gi])
            info["loop"] = {"first_row": first, "last_row": run[-1], "headers": uniq}

    def _detect_schedule(self, info):
        """Table listing tests with clause numbers -> schedule of tests (rows shown only for requested tests)."""
        rows = info["rows"]
        clause_rows = [r for r, cells in enumerate(rows) if any(CLAUSE_RE.fullmatch(c.replace("  ", " ").strip()) for c in cells)]
        if len(clause_rows) < 3:
            return
        header = rows[0]
        sheet_col = None
        for gi, h in enumerate(info_grid_headers(header)):
            if norm(h).startswith("sheet"):
                sheet_col = gi
        info["schedule"] = {"clause_rows": clause_rows, "sheet_col": sheet_col}

    def _footer_slots(self):
        for si, section in enumerate(self.doc.sections):
            for part_name, part in (("footer", section.footer), ("header", section.header)):
                root = part._element
                for pi, p in enumerate(root.iter(qn("w:p"))):
                    text = p_text(p)
                    if DUMMY_RE.search(text) or PLACEHOLDER_RE.search(text):
                        self._loc = {"part": part_name, "block": None, "sheet": 0, "fpara": pi}
                        self._spans(p, text, f"{part_name[0].upper()}{si}p{pi}", {"heading": part_name}, "para")

    # ------------------------------------------------------------------ export
    def to_json(self):
        return {"blocks": self.blocks, "sheets": self.sheets, "slots": self.slots, "tables": self.tables}


def tab_text(p):
    """Paragraph text with w:tab rendered as a space (for detection only)."""
    out = []
    for el in p.iter(qn("w:t"), qn("w:tab")):
        if next(el.iterancestors(qn("w:p")), None) is not p:
            continue
        out.append(" " if el.tag == qn("w:tab") else (el.text or ""))
    return "".join(out)


def unique_cells_text(row):
    return [(gi, c.text.strip()) for gi, c in unique_cells(row)]


def info_grid_headers(header_cells):
    return header_cells


def parse(path_or_doc, sheet_marker=r"^TEST REPORT$"):
    doc = Document(path_or_doc) if isinstance(path_or_doc, str) else path_or_doc
    return ParsedDoc(doc, sheet_marker)
